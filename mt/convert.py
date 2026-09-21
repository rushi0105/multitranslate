"""Download-time format conversion: the user uploaded an .xlsx but wants .csv (or the reverse),
a .docx but wants .txt ... Only cheap, lossless-enough conversions are offered."""
from __future__ import annotations

import csv
import io
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook

# which download formats each source extension can be turned into
CONVERSIONS: dict[str, tuple[str, ...]] = {
    ".xlsx": (".xlsx", ".csv"),
    ".csv": (".csv", ".xlsx"),
    ".tsv": (".tsv", ".xlsx"),
    ".docx": (".docx", ".txt"),
    ".txt": (".txt", ".docx"),
    ".md": (".md", ".txt"),
}
MIME = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv", ".tsv": "text/tab-separated-values", ".txt": "text/plain", ".md": "text/markdown",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".json": "application/json", ".srt": "application/x-subrip", ".vtt": "text/vtt",
}


def options_for(path: Path) -> tuple[str, ...]:
    ext = path.suffix.lower()
    return CONVERSIONS.get(ext, (ext,))


def convert(path: Path, target_ext: str) -> bytes:
    """Return the file's bytes in `target_ext` (same extension = bytes unchanged)."""
    src = path.suffix.lower()
    if target_ext == src or target_ext not in options_for(path):
        return path.read_bytes()
    if src in (".xlsx",) and target_ext == ".csv":
        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        buf = io.StringIO()
        w = csv.writer(buf)
        for row in ws.iter_rows(values_only=True):
            w.writerow(["" if v is None else v for v in row])
        return ("﻿" + buf.getvalue()).encode("utf-8")
    if src in (".csv", ".tsv") and target_ext == ".xlsx":
        text = path.read_bytes().decode("utf-8-sig", errors="replace")
        wb = Workbook()
        ws = wb.active
        for row in csv.reader(io.StringIO(text), delimiter="\t" if src == ".tsv" else ","):
            ws.append(row)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
    if src == ".docx" and target_ext == ".txt":
        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append("\t".join(c.text for c in row.cells))
        return "\n".join(lines).encode("utf-8")
    if src in (".txt", ".md") and target_ext == ".docx":
        doc = Document()
        for line in path.read_bytes().decode("utf-8-sig", errors="replace").splitlines():
            doc.add_paragraph(line)
        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()
    if src == ".md" and target_ext == ".txt":
        return path.read_bytes()
    return path.read_bytes()
