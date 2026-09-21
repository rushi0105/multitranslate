import csv
import json
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from mt.handlers import SUPPORTED_EXTENSIONS, handler_for

ROWS = [["ID", "Comment", "Rating", "Email"],
        [1, "Great soap", 5, "a@b.com"],
        [2, "Too pricey", 2, "c@d.com"]]


@pytest.fixture
def samples(tmp_path: Path) -> dict[str, Path]:
    files = {}
    (files.setdefault("txt", tmp_path / "n.txt")).write_text("Hello\n\nPrice: 199\n", encoding="utf-8")
    (files.setdefault("md", tmp_path / "n.md")).write_text(
        "# Title\n\nSome `code` here\n\n```\nkeep me\n```\n\n| Name | Age |\n|---|---|\n| Bob | 3 |\n- item one\n", encoding="utf-8")
    with (files.setdefault("csv", tmp_path / "c.csv")).open("w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(ROWS)
    wb = Workbook(); ws = wb.active; ws.title = "S1"
    for r in ROWS: ws.append(r)
    ws.append(["Total", "=SUM(C2:C3)", None, None])
    wb.save(files.setdefault("xlsx", tmp_path / "c.xlsx"))
    d = Document(); d.add_paragraph("Use daily"); t = d.add_table(rows=1, cols=1); t.rows[0].cells[0].text = "100 grams"
    d.save(files.setdefault("docx", tmp_path / "g.docx"))
    (files.setdefault("json", tmp_path / "i.json")).write_text(
        json.dumps({"greeting": "Hello {name}", "n": 3, "list": ["One", {"k": "Two"}]}), encoding="utf-8")
    (files.setdefault("srt", tmp_path / "s.srt")).write_text(
        "1\n00:00:01,000 --> 00:00:02,000\nHello there\nSecond line\n\n2\n00:00:03,000 --> 00:00:04,000\nBye\n", encoding="utf-8")
    (files.setdefault("vtt", tmp_path / "s.vtt")).write_text(
        "WEBVTT\n\n00:01.000 --> 00:02.000\nHello\n", encoding="utf-8")
    return files


def test_registry_covers_every_required_format():
    for ext in (".txt", ".md", ".csv", ".xlsx", ".docx", ".json", ".srt", ".vtt"):
        assert ext in SUPPORTED_EXTENSIONS and handler_for(Path(f"x{ext}")) is not None
    assert handler_for(Path("x.exe")) is None


def test_txt_separate_and_combined(samples, translator, tmp_path):
    h = handler_for(samples["txt"])
    h.translate_separate(samples["txt"], tmp_path / "o.txt", translator, "hi")
    assert (tmp_path / "o.txt").read_text(encoding="utf-8") == "[hi] Hello\n\n[hi] Price: 199"
    h.translate_combined(samples["txt"], tmp_path / "m.txt", translator, ["hi", "pa"], "columns")
    out = (tmp_path / "m.txt").read_text(encoding="utf-8")
    assert "===== Original =====" in out and "===== Hindi (hi) =====" in out and "[pa] Hello" in out


def test_markdown_keeps_structure(samples, translator, tmp_path):
    handler_for(samples["md"]).translate_separate(samples["md"], tmp_path / "o.md", translator, "hi")
    out = (tmp_path / "o.md").read_text(encoding="utf-8")
    assert "# [hi] Title" in out and "keep me" in out and "[hi] keep me" not in out
    assert "|---|---|" in out and "| [hi] Name | [hi] Age |" in out and "- [hi] item one" in out
    assert "`code`" in out                                  # inline code protected


def test_csv_combined_columns(samples, translator, tmp_path):
    handler_for(samples["csv"]).translate_combined(samples["csv"], tmp_path / "m.csv", translator, ["hi", "pa"], "columns")
    rows = list(csv.reader((tmp_path / "m.csv").open(encoding="utf-8-sig")))
    assert rows[0] == ["ID", "Comment", "Comment (Hindi)", "Comment (Punjabi)", "Rating", "Email"]
    assert rows[1] == ["1", "Great soap", "[hi] Great soap", "[pa] Great soap", "5", "a@b.com"]


def test_xlsx_separate_keeps_numbers_and_formulas(samples, translator, tmp_path):
    handler_for(samples["xlsx"]).translate_separate(samples["xlsx"], tmp_path / "o.xlsx", translator, "hi")
    ws = load_workbook(tmp_path / "o.xlsx")["S1"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][1] == "[hi] Comment" and rows[1] == (1, "[hi] Great soap", 5, "a@b.com")
    assert rows[3][1] == "=SUM(C2:C3)"


def test_xlsx_combined_sheets_mode(samples, translator, tmp_path):
    handler_for(samples["xlsx"]).translate_combined(samples["xlsx"], tmp_path / "m.xlsx", translator, ["hi"], "sheets")
    wb = load_workbook(tmp_path / "m.xlsx")
    assert wb.sheetnames == ["S1", "S1 (hi)"]
    assert wb["S1"]["B2"].value == "Great soap" and wb["S1 (hi)"]["B2"].value == "[hi] Great soap"


def test_xlsx_combined_columns_mode(samples, translator, tmp_path):
    handler_for(samples["xlsx"]).translate_combined(samples["xlsx"], tmp_path / "m.xlsx", translator, ["hi", "pa"], "columns")
    ws = load_workbook(tmp_path / "m.xlsx")["S1"]
    header = [c.value for c in ws[1]]
    # column A also holds the text "Total", so it counts as a text column and gets its own translations
    assert header[:6] == ["ID", "ID (Hindi)", "ID (Punjabi)", "Comment", "Comment (Hindi)", "Comment (Punjabi)"]
    assert ws["E2"].value == "[hi] Great soap"


def test_docx(samples, translator, tmp_path):
    handler_for(samples["docx"]).translate_separate(samples["docx"], tmp_path / "o.docx", translator, "hi")
    d = Document(str(tmp_path / "o.docx"))
    assert d.paragraphs[0].text == "[hi] Use daily" and d.tables[0].rows[0].cells[0].text == "[hi] 100 grams"
    handler_for(samples["docx"]).translate_combined(samples["docx"], tmp_path / "m.docx", translator, ["hi", "pa"], "columns")
    texts = [p.text for p in Document(str(tmp_path / "m.docx")).paragraphs]
    assert "Use daily" in texts and "[pa] Use daily" in texts


def test_json_values_only(samples, translator, tmp_path):
    handler_for(samples["json"]).translate_separate(samples["json"], tmp_path / "o.json", translator, "hi")
    out = json.loads((tmp_path / "o.json").read_text(encoding="utf-8"))
    assert out == {"greeting": "[hi] Hello {name}", "n": 3, "list": ["[hi] One", {"k": "[hi] Two"}]}
    handler_for(samples["json"]).translate_combined(samples["json"], tmp_path / "m.json", translator, ["hi"], "columns")
    multi = json.loads((tmp_path / "m.json").read_text(encoding="utf-8"))
    assert set(multi) == {"original", "hi"} and multi["original"]["n"] == 3


def test_srt_and_vtt(samples, translator, tmp_path):
    handler_for(samples["srt"]).translate_separate(samples["srt"], tmp_path / "o.srt", translator, "hi")
    out = (tmp_path / "o.srt").read_text(encoding="utf-8")
    assert "00:00:01,000 --> 00:00:02,000\n[hi] Hello there\n[hi] Second line\n\n2\n" in out
    handler_for(samples["srt"]).translate_combined(samples["srt"], tmp_path / "m.srt", translator, ["hi", "pa"], "columns")
    assert "Bye\n[hi] Bye\n[pa] Bye" in (tmp_path / "m.srt").read_text(encoding="utf-8")
    handler_for(samples["vtt"]).translate_separate(samples["vtt"], tmp_path / "o.vtt", translator, "hi")
    assert (tmp_path / "o.vtt").read_text(encoding="utf-8").startswith("WEBVTT\n\n00:01.000 --> 00:02.000\n[hi] Hello")
